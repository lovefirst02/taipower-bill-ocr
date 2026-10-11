import tempfile
import unittest
from pathlib import Path

from taipower_bill_ocr.excel_export import (
    COLUMNS,
    bill_to_record,
    bill_to_row,
    bills_to_rows,
    write_rows_to_excel,
)
from taipower_bill_ocr.parser import BillResult


def high_voltage_bill(**charge_overrides):
    charge_info = {
        "basic_charge_twd": 700000,
        "energy_charge_twd": 120000,
        "power_factor_adjustment_twd": -25000,
        "pretax_total_twd": 761905,
        "business_tax_twd": 38095,
        "total_amount_twd": 800000,
        **charge_overrides,
    }
    charge_info = {key: value for key, value in charge_info.items() if value is not None}
    return BillResult(
        {
            "user_info": {
                "customer_number": "00-12-3456-78-9",
                "billing_period": {"start": "109/12/01", "end": "109/12/31"},
                "energy_consumption": {
                    "peak_kwh": 8000,
                    "semi_peak_kwh": 3000,
                    "saturday_semi_peak_kwh": 3000,
                    "off_peak_kwh": 1000,
                },
                "transferred_energy": {
                    "peak_kwh": 100,
                    "semi_peak_kwh": 200,
                    "saturday_semi_peak_kwh": 300,
                    "off_peak_kwh": 400,
                },
            },
            "charge_info": charge_info,
            "quality": {"warnings": [], "page_number": 1},
        }
    )


class ExcelRowTests(unittest.TestCase):
    def test_maps_every_requested_column(self):
        row = bill_to_row(high_voltage_bill(prepaid_applied_twd=-5000), "bill.pdf")

        self.assertEqual(row["Billing Period 計費期間"], "109/12/01~109/12/31")
        self.assertEqual(row["Custom No. 電號"], "00-12-3456-78-9")
        self.assertEqual(row["計費度數_尖峰度數"], 8000)
        self.assertEqual(row["計費度數_週六半尖峰度數"], 3000)
        self.assertEqual(row["轉供度數_半尖峰度數"], 200)
        self.assertEqual(row["轉供度數_離峰度數"], 400)
        self.assertEqual(row["基本電費"], 700000)
        self.assertEqual(row["流動電費"], 120000)
        self.assertEqual(row["功率因數調整費"], -25000)
        self.assertEqual(row["預繳費用"], -5000)
        self.assertEqual(row["稅前應繳總金額"], 761905)
        self.assertEqual(row["應繳總金額"], 800000)
        self.assertEqual(row["來源檔案"], "bill.pdf")
        self.assertEqual(row["頁碼"], 1)

    def test_other_adjustment_is_pretax_minus_known_charges(self):
        row = bill_to_row(high_voltage_bill())
        # 761905 - 700000 - 120000 - (-25000)
        self.assertEqual(row["其他調整費用"], -33095)
        self.assertIsNone(row["需人工檢查"])

    def test_detailed_layout_reads_energy_usage_block(self):
        bill = high_voltage_bill()
        bill["energy_usage"] = bill["user_info"].pop("energy_consumption")
        bill["energy_usage"] = {"energy_consumption": bill["energy_usage"]}
        self.assertEqual(bill_to_row(bill)["計費度數_離峰度數"], 1000)

    def test_missing_inputs_leave_other_adjustment_empty_and_flagged(self):
        row = bill_to_row(high_voltage_bill(energy_charge_twd=None))
        self.assertIsNone(row["其他調整費用"])
        self.assertIn("缺少流動電費", row["需人工檢查"])

    def test_missing_power_factor_counts_as_zero_and_is_flagged(self):
        row = bill_to_row(
            high_voltage_bill(power_factor_adjustment_twd=None, pretax_total_twd=830000,
                              total_amount_twd=871500, business_tax_twd=41500)
        )
        self.assertEqual(row["其他調整費用"], 10000)
        self.assertIn("功率因數調整費", row["需人工檢查"])

    def test_flags_reversed_billing_period(self):
        bill = high_voltage_bill()
        bill["user_info"]["billing_period"] = {"start": "113/04/30", "end": "113/04/01"}
        self.assertIn("計費期間起訖異常", bill_to_row(bill)["需人工檢查"])

    def test_flags_total_inconsistent_with_pretax(self):
        row = bill_to_row(high_voltage_bill(total_amount_twd=120000))
        self.assertIn("應繳總金額與稅前金額", row["需人工檢查"])

    def test_misread_tax_does_not_flag_consistent_total(self):
        row = bill_to_row(high_voltage_bill(business_tax_twd=1710))
        self.assertIsNone(row["需人工檢查"])

    def test_skips_pages_without_charge_breakdown(self):
        back_page = BillResult({"meter_info": {}, "quality": {"warnings": []}})
        rows = bills_to_rows([high_voltage_bill(), back_page])
        self.assertEqual(len(rows), 1)


class ExcelWriteTests(unittest.TestCase):
    def test_writes_header_and_keeps_customer_number_as_text(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "bills.xlsx"
            write_rows_to_excel([bill_to_row(high_voltage_bill())], path)
            sheet = load_workbook(path).active

            self.assertEqual([cell.value for cell in sheet[1]], list(COLUMNS))
            customer_cell = sheet.cell(row=2, column=2)
            self.assertEqual(customer_cell.value, "00-12-3456-78-9")
            self.assertEqual(customer_cell.number_format, "@")
            self.assertEqual(sheet.cell(row=2, column=15).value, -33095)



class CrossCheckSuppressionTests(unittest.TestCase):
    def test_reversed_period_from_cross_check_is_not_flagged(self):
        other = high_voltage_bill()
        other["user_info"]["billing_period"] = {"start": "109/12/31", "end": "109/12/01"}
        record = bill_to_record(high_voltage_bill(), cross_check=other)
        self.assertEqual(record.issues, [])
        self.assertEqual(record.alternatives, {})

    def test_total_inconsistent_with_pretax_is_not_flagged(self):
        record = bill_to_record(
            high_voltage_bill(), cross_check=high_voltage_bill(total_amount_twd=120000)
        )
        self.assertEqual(record.issues, [])

    def test_disagreement_without_arithmetic_check_is_flagged(self):
        other = high_voltage_bill()
        other["user_info"]["energy_consumption"]["peak_kwh"] = 80
        record = bill_to_record(high_voltage_bill(), cross_check=other)
        self.assertEqual(record.alternatives, {"計費度數_尖峰度數": 80})

    def test_flagged_when_primary_fails_the_check(self):
        record = bill_to_record(
            high_voltage_bill(total_amount_twd=120000),
            cross_check=high_voltage_bill(),
        )
        self.assertEqual(record.alternatives, {"應繳總金額": 800000})


if __name__ == "__main__":
    unittest.main()
