import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from taipower_bill_ocr import watcher
from taipower_bill_ocr.watcher import FolderProcessor, SingleInstanceLock, WatchConfig

from test_batch_gui import FakeExtractor


class FakeClock:
    def __init__(self, now):
        self.now = now

    def __call__(self):
        return self.now


class FolderProcessorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.inbox = root / "Bills"
        self.output = root / "Bills OCR"
        self.inbox.mkdir()
        self.clock = FakeClock(1_000_000.0)
        self.extractor = FakeExtractor()
        self.config = WatchConfig(self.inbox, self.output, settle_seconds=30)

    def tearDown(self):
        self.temp.cleanup()

    def add(self, name, content=b"bill", age=60):
        path = self.inbox / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        mtime = self.clock.now - age
        os.utime(path, (mtime, mtime))
        return path

    def processor(self):
        return FolderProcessor(self.config, extractor=self.extractor, clock=self.clock)

    def sheet_rows(self, name=None):
        workbook = load_workbook(self.output / watcher.WORKBOOK)
        sheet = workbook[name] if name else workbook.active
        return [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]

    def test_new_files_become_rows_in_the_cumulative_workbook(self):
        self.add("2024/jan.png")
        self.add("feb.png")
        summary = self.processor().run_once()

        self.assertEqual(summary, {"processed": 2, "failed": 0})
        sources = [row[17] for row in self.sheet_rows()]
        self.assertEqual(sorted(sources), ["2024/jan.png", "feb.png"])
        files = self.sheet_rows("Files")
        self.assertEqual([(row[0], row[1], row[2]) for row in files], [("2024/jan.png", "done", 1), ("feb.png", "done", 1)])

    def test_unchanged_files_are_not_processed_again(self):
        self.add("jan.png")
        self.processor().run_once()
        calls = len(self.extractor.calls)

        summary = self.processor().run_once()  # a fresh run reloads the saved state
        self.assertEqual(summary["processed"], 0)
        self.assertEqual(len(self.extractor.calls), calls)

    def test_file_still_syncing_waits_until_it_settles(self):
        self.add("jan.png", age=5)
        self.assertEqual(self.processor().run_once()["processed"], 0)
        self.clock.now += 60
        self.assertEqual(self.processor().run_once()["processed"], 1)

    def test_changed_content_is_reprocessed_and_replaces_its_rows(self):
        path = self.add("jan.png", b"v1")
        self.processor().run_once()
        path.write_bytes(b"version 2")
        os.utime(path, (self.clock.now - 40, self.clock.now - 40))

        self.assertEqual(self.processor().run_once()["processed"], 1)
        self.assertEqual(len(self.sheet_rows()), 1)  # replaced, not duplicated

    def test_touched_but_identical_file_is_not_reprocessed(self):
        path = self.add("jan.png")
        self.processor().run_once()
        calls = len(self.extractor.calls)
        os.utime(path, (self.clock.now - 35, self.clock.now - 35))  # e.g. re-synced

        self.processor().run_once()
        self.assertEqual(len(self.extractor.calls), calls)

    def test_failures_are_retried_up_to_the_limit(self):
        self.add("broken.png")
        self.extractor = FakeExtractor(fail_on="broken.png")
        for _ in range(5):
            self.processor().run_once()

        self.assertEqual(len(self.extractor.calls), self.config.max_attempts)
        entry = json.loads((self.output / watcher.STATE_FILE).read_text(encoding="utf-8"))["files"]["broken.png"]
        self.assertEqual((entry["status"], entry["attempts"]), ("failed", 3))
        self.assertIn("unreadable", self.sheet_rows("Files")[0][6])

    def test_temporary_and_output_files_are_ignored(self):
        self.add("~$draft.png")
        self.config = WatchConfig(self.inbox, self.inbox / "out", settle_seconds=30)
        self.add("out/old_page.png")
        self.assertEqual(self.processor().pending_files(), [])

    def test_workbook_locked_by_excel_is_retried_next_scan(self):
        self.add("jan.png")
        processor = self.processor()
        real_replace = os.replace

        def locked(source, target):
            if str(target).endswith(watcher.WORKBOOK):
                raise PermissionError("in use")
            return real_replace(source, target)

        with mock.patch.object(watcher.os, "replace", side_effect=locked):
            processor.run_once()
        self.assertTrue(processor.workbook_dirty)
        self.assertFalse((self.output / watcher.WORKBOOK).exists())
        self.assertEqual([p.name for p in self.output.iterdir() if p.name.startswith("~")], [])

        processor.run_once()
        self.assertFalse(processor.workbook_dirty)
        self.assertEqual(len(self.sheet_rows()), 1)


class SingleInstanceLockTests(unittest.TestCase):
    def test_second_watcher_on_same_output_is_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            with SingleInstanceLock(Path(folder)):
                with self.assertRaises(RuntimeError):
                    with SingleInstanceLock(Path(folder)):
                        pass
            with SingleInstanceLock(Path(folder)):  # released after the first exits
                pass


if __name__ == "__main__":
    unittest.main()
