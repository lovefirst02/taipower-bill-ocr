"""The master workbook: everything read from bills as Power Query tables.

Shared by the folder watcher and the web app so both write the same format.
A *document* is the per-bill JSON written by either: file, processed_at,
program_version and its pages (each a ``PageRecord.as_json()``).
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

PAGE_COLUMNS = [
    "file", "page", "document_type", "customer_number", "billing_period_start",
    "billing_period_end", "total_amount_twd", "needs_review", "fields_recognized",
    "fields_unrecognized", "engines_agree", "engines_disagree",
    "filled_from_cross_check", "primary_only", "unverified", "resolved_by_check",
    "confidence_high", "confidence_medium", "confidence_low", "confidence_review",
    "charges_difference_twd", "tax_difference_twd", "supplemented_from_ocr_lines", "warnings",
]
_CONTEXT = ["file", "page", "customer_number", "billing_period_start", "billing_period_end"]
SHEET_COLUMNS = {
    "Pages": [*PAGE_COLUMNS, "processed_at", "program_version"],
    "Fields": [
        *_CONTEXT, "block", "block_label", "key", "label", "printed_label", "value",
        "recognized", "status", "confidence", "checks", "cross_check_value", "source", "origin",
        "x", "y", "width", "height",
    ],
    "TableRows": [
        *_CONTEXT, "block", "table", "table_label", "row", "column", "column_label", "value",
    ],
    "Files": ["file", "status", "pages", "needs_review", "processed_at", "attempts", "error"],
}


def empty_sheets() -> dict[str, list[list[Any]]]:
    return {name: [] for name in SHEET_COLUMNS}


def add_document(sheets: dict[str, list[list[Any]]], document: Mapping[str, Any]) -> None:
    """Append one bill's pages, fields and table cells to ``sheets``."""
    for page in document["pages"]:
        summary = page["summary"]
        context = [summary["file"], summary["page"], summary["customer_number"],
                   summary["billing_period_start"], summary["billing_period_end"]]
        sheets["Pages"].append(
            [summary.get(column) for column in PAGE_COLUMNS]
            + [document["processed_at"], document["program_version"]]
        )
        for item in page["fields"]:
            box = item["bbox"] or {}
            sheets["Fields"].append(
                context
                + [item["block"], item["block_label"], item["key"], item["label"],
                   item.get("printed_label"), item["value"], item["recognized"],
                   item["status"], item.get("confidence"), item.get("checks"),
                   item["cross_check_value"], item["source"], item.get("origin"),
                   box.get("x"), box.get("y"), box.get("width"), box.get("height")]
            )
        for cell in page["table_rows"]:
            sheets["TableRows"].append(
                context
                + [cell["block"], cell["table"], cell["table_label"], cell["row"],
                   cell["column"], cell["column_label"], cell["value"]]
            )


def save_tables(
    target: Path,
    tables: Mapping[str, tuple[Sequence[str], Iterable[Sequence[Any]]]],
) -> bool:
    """Write each ``name -> (columns, rows)`` as a sheet holding an Excel table.

    The file is replaced atomically; returns False (leaving the old file) when
    another program, typically Excel, has it open.
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, (columns, rows) in tables.items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(columns))
        count = 0
        for row in rows:
            sheet.append(list(row))
            count += 1
        sheet.freeze_panes = "A2"
        for index, column in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(12, min(len(str(column)) + 4, 40))
        if count:
            table = Table(displayName=name, ref=f"A1:{get_column_letter(len(columns))}{count + 1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)

    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_name(f"~tmp_{target.name}")
    workbook.save(temp)
    try:
        os.replace(temp, target)
    except PermissionError:
        temp.unlink(missing_ok=True)
        return False
    return True


def save_master(target: Path, sheets: Mapping[str, list[list[Any]]]) -> bool:
    return save_tables(target, {name: (SHEET_COLUMNS[name], sheets[name]) for name in SHEET_COLUMNS})
