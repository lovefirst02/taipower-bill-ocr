"""Watch a folder (e.g. a OneDrive-synced SharePoint library) and OCR new bills.

Everything read from each new or changed PDF/image is kept: one JSON file
per bill under ``json/`` (including raw OCR lines), and a cumulative
workbook whose Pages, Fields, TableRows and Files tables consumers query for
whatever they need. Run continuously, or once per call from Task Scheduler::

    python -m taipower_bill_ocr.watcher "C:\\Users\\me\\Company\\Bills" --output "C:\\Users\\me\\Company\\Bills OCR"
    python -m taipower_bill_ocr.watcher IN --output OUT --once

Polling is used instead of file-system events because OneDrive writes files
in several steps; a file is only read once it has stopped changing.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import sys
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from .batch import FileFinished, PageDone, PageExtractor, find_bill_files, iter_batch
from .dataset import PROGRAM_VERSION, SCHEMA_VERSION, page_record

STATE_FILE = ".taipower_ocr_state.json"
LOCK_FILE = ".taipower_ocr.lock"
LOG_FILE = "taipower_ocr.log"
WORKBOOK = "taipower_bills.xlsx"
JSON_DIR = "json"

PAGE_COLUMNS = [
    "file", "page", "document_type", "customer_number", "billing_period_start",
    "billing_period_end", "total_amount_twd", "needs_review", "fields_recognized",
    "fields_unrecognized", "engines_agree", "engines_disagree",
    "filled_from_cross_check", "primary_only", "warnings",
]
_CONTEXT = ["file", "page", "customer_number", "billing_period_start", "billing_period_end"]
SHEET_COLUMNS = {
    "Pages": [*PAGE_COLUMNS, "processed_at", "program_version"],
    "Fields": [
        *_CONTEXT, "block", "block_label", "key", "label", "value", "recognized",
        "status", "cross_check_value", "source", "x", "y", "width", "height",
    ],
    "TableRows": [
        *_CONTEXT, "block", "table", "table_label", "row", "column", "column_label", "value",
    ],
    "Files": ["file", "status", "pages", "needs_review", "processed_at", "attempts", "error"],
}

log = logging.getLogger("taipower_bill_ocr.watcher")


@dataclass
class WatchConfig:
    input_dir: Path
    output_dir: Path
    recursive: bool = True
    interval: float = 60.0
    # A file must be unchanged this long before it is read, so a file OneDrive
    # is still downloading is not processed half-written.
    settle_seconds: float = 30.0
    max_attempts: int = 3
    ocr_backend: str = "crosscheck"
    dpi: int = 200


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _is_temporary(path: Path) -> bool:
    name = path.name
    return name.startswith(("~$", ".")) or name.lower().endswith((".tmp", ".partial"))


class FolderProcessor:
    """Keeps a JSON record of processed files and the rows they produced."""

    def __init__(
        self,
        config: WatchConfig,
        extractor: Any | None = None,
        clock: Callable[[], float] = time.time,
    ) -> None:
        self.config = config
        self.clock = clock
        self._extractor = extractor
        config.output_dir.mkdir(parents=True, exist_ok=True)
        self.state_path = config.output_dir / STATE_FILE
        self.state: dict[str, Any] = self._load_state()
        self.workbook_dirty = False

    # ------------------------------------------------------------------ state
    def _load_state(self) -> dict[str, Any]:
        try:
            return json.loads(self.state_path.read_text(encoding="utf-8"))
        except FileNotFoundError:
            return {"files": {}}

    def _save_state(self) -> None:
        temp = self.state_path.with_suffix(".tmp")
        temp.write_text(json.dumps(self.state, ensure_ascii=False, indent=1), encoding="utf-8")
        os.replace(temp, self.state_path)

    @property
    def extractor(self) -> Any:
        if self._extractor is None:
            log.info("Loading OCR engine (%s)…", self.config.ocr_backend)
            self._extractor = PageExtractor(self.config.ocr_backend, self.config.dpi)
            if self._extractor.uses_paddle:
                self._extractor._engine()
        return self._extractor

    # ------------------------------------------------------------------- scan
    def _key(self, path: Path) -> str:
        return path.relative_to(self.config.input_dir).as_posix()

    def pending_files(self) -> list[Path]:
        """Settled files that are new, changed, or failed fewer than max times."""
        output = self.config.output_dir.resolve()
        pending = []
        for path in find_bill_files(self.config.input_dir, recursive=self.config.recursive):
            if _is_temporary(path) or output in path.resolve().parents:
                continue
            try:
                stat = path.stat()
            except OSError:
                continue  # removed or being replaced by the sync client
            if self.clock() - stat.st_mtime < self.config.settle_seconds:
                continue
            entry = self.state["files"].get(self._key(path))
            if entry and entry["size"] == stat.st_size and entry["mtime"] == stat.st_mtime:
                if entry["status"] == "done" or entry["attempts"] >= self.config.max_attempts:
                    continue
            pending.append(path)
        return pending

    # ---------------------------------------------------------------- process
    def process(self, path: Path) -> None:
        key = self._key(path)
        stat = path.stat()
        digest = _sha256(path)
        previous = self.state["files"].get(key, {})
        if previous.get("sha256") == digest and previous.get("status") == "done":
            # Same content with a new timestamp (e.g. re-synced): nothing to redo.
            previous.update(size=stat.st_size, mtime=stat.st_mtime)
            self._save_state()
            return
        attempts = previous.get("attempts", 0) + 1 if previous.get("sha256") == digest else 1

        pages: list[dict[str, Any]] = []
        errors: list[str] = []
        started = time.perf_counter()
        for event in iter_batch(
            [path], self.extractor, source_name=lambda _: key, cross_check_all_pages=True
        ):
            if isinstance(event, PageDone) and event.bill is not None:
                record = page_record(
                    event.bill,
                    file=key,
                    page=event.page_number,
                    cross_check=event.cross_check,
                    cross_check_error=event.cross_check_error,
                )
                pages.append(record.as_json())
            elif isinstance(event, FileFinished) and event.error:
                errors.append(event.error)

        entry = {
            "sha256": digest,
            "size": stat.st_size,
            "mtime": stat.st_mtime,
            "processed_at": datetime.now().isoformat(timespec="seconds"),
            "attempts": attempts,
        }
        if errors:
            entry.update(status="failed", error="; ".join(errors))
            log.warning("Failed %s (attempt %d/%d): %s", key, attempts, self.config.max_attempts, entry["error"])
        else:
            review = sum(1 for page in pages if page["summary"]["needs_review"])
            entry.update(status="done", pages=len(pages), needs_review=review)
            self._write_json(key, digest, entry["processed_at"], pages)
            log.info("Processed %s: %d page(s), %d need review, %.0fs", key, len(pages), review, time.perf_counter() - started)
        self.state["files"][key] = entry
        self._save_state()
        self.workbook_dirty = True

    def _json_path(self, key: str) -> Path:
        return self.config.output_dir / JSON_DIR / f"{key}.json"

    def _write_json(self, key: str, digest: str, processed_at: str, pages: list[dict[str, Any]]) -> None:
        document = {
            "schema_version": SCHEMA_VERSION,
            "program_version": PROGRAM_VERSION,
            "file": key,
            "sha256": digest,
            "processed_at": processed_at,
            "ocr": {
                "primary": "windows" if self.config.ocr_backend == "windows" else "paddleocr",
                "cross_check": "windows" if self.config.ocr_backend == "crosscheck" else None,
                "dpi": self.config.dpi,
            },
            "pages": pages,
        }
        target = self._json_path(key)
        target.parent.mkdir(parents=True, exist_ok=True)
        temp = target.with_suffix(".tmp")
        temp.write_text(json.dumps(document, ensure_ascii=False, indent=1), encoding="utf-8")
        os.replace(temp, target)

    # --------------------------------------------------------------- workbook
    def _sheet_rows(self) -> dict[str, list[list[Any]]]:
        sheets: dict[str, list[list[Any]]] = {name: [] for name in SHEET_COLUMNS}
        for key, entry in sorted(self.state["files"].items()):
            sheets["Files"].append(
                [key, entry["status"], entry.get("pages"), entry.get("needs_review"),
                 entry["processed_at"], entry["attempts"], entry.get("error")]
            )
            if entry["status"] != "done":
                continue
            try:
                document = json.loads(self._json_path(key).read_text(encoding="utf-8"))
            except FileNotFoundError:
                continue
            for page in document["pages"]:
                summary = page["summary"]
                context = [summary["file"], summary["page"], summary["customer_number"],
                           summary["billing_period_start"], summary["billing_period_end"]]
                sheets["Pages"].append(
                    [summary.get(column) for column in PAGE_COLUMNS]
                    + [document["processed_at"], document["program_version"]]
                )
                for item in page["fields"]:
                    box = item["bbox"] or {}
                    sheets["Fields"].append(
                        context
                        + [item["block"], item["block_label"], item["key"], item["label"],
                           item["value"], item["recognized"], item["status"],
                           item["cross_check_value"], item["source"],
                           box.get("x"), box.get("y"), box.get("width"), box.get("height")]
                    )
                for cell in page["table_rows"]:
                    sheets["TableRows"].append(
                        context
                        + [cell["block"], cell["table"], cell["table_label"], cell["row"],
                           cell["column"], cell["column_label"], cell["value"]]
                    )
        return sheets

    def write_workbook(self) -> bool:
        """Rewrite the cumulative workbook; False if Excel has it locked.

        Every sheet is an Excel table named like the sheet, so Power Query can
        read Pages, Fields, TableRows and Files directly.
        """
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        sheets = self._sheet_rows()
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, columns in SHEET_COLUMNS.items():
            sheet = workbook.create_sheet(name)
            sheet.append(columns)
            for row in sheets[name]:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            for index, column in enumerate(columns, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = max(12, min(len(column) + 4, 40))
            if sheets[name]:
                table = Table(displayName=name, ref=f"A1:{get_column_letter(len(columns))}{len(sheets[name]) + 1}")
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                sheet.add_table(table)

        target = self.config.output_dir / WORKBOOK
        temp = target.with_name(f"~tmp_{target.name}")
        workbook.save(temp)
        try:
            os.replace(temp, target)
        except PermissionError:
            log.warning("%s is open in another program; will retry next scan", target.name)
            temp.unlink(missing_ok=True)
            return False
        self.workbook_dirty = False
        log.info(
            "Updated %s (%d page(s), %d field(s))",
            target.name, len(sheets["Pages"]), len(sheets["Fields"]),
        )
        return True

    def run_once(self) -> dict[str, int]:
        pending = self.pending_files()
        for path in pending:
            try:
                self.process(path)
            except OSError as exc:  # vanished or locked mid-sync; next scan retries
                log.warning("Skipped %s for now: %s", path.name, exc)
        if self.workbook_dirty:
            self.write_workbook()
        failed = sum(1 for entry in self.state["files"].values() if entry["status"] == "failed")
        return {"processed": len(pending), "failed": failed}


class SingleInstanceLock:
    """Stops two watchers from writing the same output folder."""

    def __init__(self, folder: Path) -> None:
        self.path = folder / LOCK_FILE
        self.handle: Any = None

    def __enter__(self) -> "SingleInstanceLock":
        import msvcrt

        self.handle = open(self.path, "a+")  # noqa: SIM115
        self.handle.seek(0)  # every instance must contend for the same byte
        try:
            msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
        except OSError:
            self.handle.close()
            raise RuntimeError(f"Another watcher is already using {self.path.parent}") from None
        return self

    def __exit__(self, *exc: object) -> None:
        import msvcrt

        self.handle.seek(0)
        msvcrt.locking(self.handle.fileno(), msvcrt.LK_UNLCK, 1)
        self.handle.close()


def watch(config: WatchConfig, *, once: bool = False, stop: threading.Event | None = None) -> int:
    stop = stop or threading.Event()
    with SingleInstanceLock(config.output_dir):
        processor = FolderProcessor(config)
        log.info("Watching %s -> %s", config.input_dir, config.output_dir)
        while True:
            summary = processor.run_once()
            if once:
                return 1 if summary["failed"] else 0
            if stop.wait(config.interval):
                return 0


def _configure_logging(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s", "%Y-%m-%d %H:%M:%S")
    handlers: list[logging.Handler] = [logging.FileHandler(output_dir / LOG_FILE, encoding="utf-8")]
    if sys.stderr is not None:
        handlers.append(logging.StreamHandler())
    for handler in handlers:
        handler.setFormatter(formatter)
        log.addHandler(handler)
    log.setLevel(logging.INFO)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Watch a folder for Taipower bills and keep an Excel summary up to date."
    )
    parser.add_argument("input", type=Path, help="Folder to watch, e.g. a OneDrive-synced SharePoint folder")
    parser.add_argument("--output", type=Path, required=True, help="Folder for taipower_bills.xlsx and the log")
    parser.add_argument("--once", action="store_true", help="Process pending files once and exit (for Task Scheduler)")
    parser.add_argument("--interval", type=float, default=60, help="Seconds between scans (default 60)")
    parser.add_argument("--settle", type=float, default=30, help="Seconds a file must be unchanged before reading (default 30)")
    parser.add_argument("--no-subfolders", action="store_true", help="Only watch the top folder")
    parser.add_argument("--ocr-backend", choices=("crosscheck", "paddleocr", "windows"), default="crosscheck")
    parser.add_argument("--dpi", type=int, default=200)
    args = parser.parse_args(argv)
    if not args.input.is_dir():
        parser.error(f"input folder not found: {args.input}")

    config = WatchConfig(
        input_dir=args.input.resolve(),
        output_dir=args.output.resolve(),
        recursive=not args.no_subfolders,
        interval=args.interval,
        settle_seconds=args.settle,
        ocr_backend=args.ocr_backend,
        dpi=args.dpi,
    )
    _configure_logging(config.output_dir)
    try:
        return watch(config, once=args.once)
    except KeyboardInterrupt:
        log.info("Stopped")
        return 0
    except RuntimeError as exc:
        log.error("%s", exc)
        return 2


if __name__ == "__main__":
    sys.exit(main())
