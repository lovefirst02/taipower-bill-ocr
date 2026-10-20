"""Budgets per month, entity and site, imported from Excel, and green power contracts.

The template has a row per month and site (with its entity); each budget
is grey, solar and wind kWh and cost, and the totals (left empty: the sum).
A green cost left empty is estimated from the contract: kWh × (the
contract's price + Taipower's transfer fees + certificate fees charged on
top), and said so.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any, Iterable, Mapping

from .mix import SOLAR, WIND, energy_name
from .store import Store

# Template column → field, in order.
COLUMNS: list[tuple[str, str]] = [
    ("月份", "month"), ("Entity", "entity"), ("Site", "site"),
    ("灰電度數", "grey_kwh"), ("灰電費用（稅前）", "grey_cost"),
    ("太陽能度數", "solar_kwh"), ("太陽能費用（未稅）", "solar_cost"),
    ("風力度數", "wind_kwh"), ("風力費用（未稅）", "wind_cost"),
    ("總用電度數", "total_kwh"), ("總費用", "total_cost"),
]
FIELDS = [field for _, field in COLUMNS[3:]]
ENERGY_FIELDS = {SOLAR: ("solar_kwh", "solar_cost"), WIND: ("wind_kwh", "wind_cost")}


def template_workbook(store: Store, year: int, existing: Iterable[Mapping[str, Any]] = ()) -> bytes:
    """A row per month and site of the year, with what is already budgeted."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    entities = {entity["id"]: entity["name"] for entity in store.entities()}
    budgeted = {(row["month"], row["site"]): row for row in existing}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append([name for name, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for month in range(1, 13):
        key = f"{year:04d}-{month:02d}"
        for site in store.sites():
            row = budgeted.get((key, site["name"]), {})
            sheet.append([key, entities.get(site.get("entity_id")), site["name"], *(row.get(field) for field in FIELDS)])
    for index, (name, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = max(12, len(name) * 2 + 2)
    sheet.freeze_panes = "D2"
    notes = workbook.create_sheet("說明")
    for line in (
        "每列：一個月份、一個 Site（廠區或辦公室）；Entity 與 Site 名稱請與系統 Settings 相同。",
        "金額：灰電為台電帳單稅前金額；綠電為售電業者未稅金額。",
        "綠電費用留空：依 Settings → Green power contracts 的合約單價推估（度數 × (合約價 + 轉供費率 + 憑證費率)）。",
        "總用電度數、總費用留空：自動加總灰電與綠電。",
        "匯入時，檔案中出現的月份會整月取代；其他月份不變。",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 100
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _month(value: Any) -> str | None:
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}-{value.month:02d}"
    match = re.fullmatch(r"\s*(\d{3,4})\s*[-/.年]\s*(\d{1,2})\s*月?\s*", str(value or ""))
    if not match:
        return None
    year, month = int(match.group(1)), int(match.group(2))
    year = year + 1911 if year < 1000 else year
    return f"{year:04d}-{month:02d}" if 1 <= month <= 12 else None


def _amount(value: Any) -> float | int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value) if float(value).is_integer() else float(value)
    text = re.sub(r"[,，\s]", "", str(value))
    try:
        number = float(text)
    except ValueError:
        raise ValueError(f"not a number: {value}") from None
    return int(number) if number.is_integer() else number


def read_workbook(content: bytes, store: Store) -> tuple[list[dict[str, Any]], list[str]]:
    """The budget rows of a filled-in template, and what is wrong with it (nothing is imported then)."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - any unreadable file
        return [], ["not an Excel file (.xlsx)"]
    sheet = workbook["Budget"] if "Budget" in workbook.sheetnames else workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], ["the sheet is empty"]
    header = [str(cell or "").strip() for cell in rows[0]]
    positions = {}
    for name, field in COLUMNS:
        if name not in header:
            return [], [f"column {name} not found (use the template)"]
        positions[field] = header.index(name)
    sites = {site["name"]: site for site in store.sites()}
    entities = {entity["id"]: entity["name"] for entity in store.entities()}
    result, errors, seen = [], [], set()
    for number, cells in enumerate(rows[1:], start=2):
        get = lambda field: cells[positions[field]] if positions[field] < len(cells) else None  # noqa: E731
        if all(get(field) in (None, "") for field in FIELDS):
            continue  # nothing budgeted on this row
        month = _month(get("month"))
        site = str(get("site") or "").strip()
        if month is None:
            errors.append(f"row {number}: month {get('month')!r} — use 2026-01")
            continue
        if site not in sites:
            errors.append(f"row {number}: site {site!r} is not in Settings → Sites")
            continue
        if (month, site) in seen:
            errors.append(f"row {number}: {month} {site} appears twice")
            continue
        seen.add((month, site))
        values: dict[str, Any] = {"month": month, "site": site, "entity": entities.get(sites[site].get("entity_id"))}
        try:
            for field in FIELDS:
                values[field] = _amount(get(field))
        except ValueError as error:
            errors.append(f"row {number}: {error}")
            continue
        result.append(values)
    if not result and not errors:
        errors.append("no budget figures found")
    return result, errors


def contract_for(contracts: Iterable[Mapping[str, Any]], entity_id: Any, energy: str, month: str) -> Mapping[str, Any] | None:
    """The contract in force for an entity's energy in a month (an entity's own before one for all)."""
    matching = [
        contract for contract in contracts
        if energy_name(contract.get("energy")) == energy
        and contract.get("entity_id") in (None, entity_id)
        and (not contract.get("start_month") or contract["start_month"] <= month)
        and (not contract.get("end_month") or month <= contract["end_month"])
    ]
    matching.sort(key=lambda contract: (contract.get("entity_id") is None, -(int(str(contract.get("start_month") or "0000-00").replace("-", "")))))
    return matching[0] if matching else None


def per_kwh(contract: Mapping[str, Any]) -> float:
    """What a kWh costs under a contract: its price and the fees charged on top."""
    return round(float(contract.get("price") or 0) + float(contract.get("transfer_rate") or 0) + float(contract.get("certificate_rate") or 0), 6)


def with_estimates(rows: Iterable[Mapping[str, Any]], store: Store) -> list[dict[str, Any]]:
    """Budgets as compared: green costs left empty estimated from the contracts, totals added up."""
    contracts = store.contracts()
    entity_ids = {entity["name"]: entity["id"] for entity in store.entities()}
    result = []
    for row in rows:
        item = dict(row)
        estimated = []
        for energy, (kwh_field, cost_field) in ENERGY_FIELDS.items():
            if item.get(cost_field) is None and item.get(kwh_field):
                contract = contract_for(contracts, entity_ids.get(item.get("entity")), energy, item["month"])
                if contract is not None:
                    item[cost_field] = round(item[kwh_field] * per_kwh(contract))
                    estimated.append(cost_field)
        parts_kwh = [item.get(field) for field in ("grey_kwh", "solar_kwh", "wind_kwh")]
        parts_cost = [item.get(field) for field in ("grey_cost", "solar_cost", "wind_cost")]
        if item.get("total_kwh") is None and any(value is not None for value in parts_kwh):
            item["total_kwh"] = sum(value or 0 for value in parts_kwh)
        if item.get("total_cost") is None and any(value is not None for value in parts_cost):
            item["total_cost"] = sum(value or 0 for value in parts_cost)
        item["estimated"] = estimated
        result.append(item)
    return result
