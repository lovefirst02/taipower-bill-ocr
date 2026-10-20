import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from taipower_bill_ocr.parser import BillResult, OcrLine
from taipower_web_api import reports
from taipower_web_api.main import create_app
from taipower_web_api.settings import Settings


def official_sample_bill():
    """Values and label positions of Taipower's official high-voltage example."""
    bill = BillResult({
        "document_type": "high_voltage_bill",
        "quality": {"warnings": []},
        "user_info": {
            "customer_number": "00-12-3456-78-9",
            "billing_period": {"start": "109/12/01", "end": "109/12/31"},
            "energy_consumption": {"peak_kwh": 8000, "semi_peak_kwh": 3000, "saturday_semi_peak_kwh": 3000, "off_peak_kwh": 1000},
        },
        "charge_info": {
            "basic_charge_twd": 700000,
            "energy_charge_twd": 120000,
            "power_factor_adjustment_twd": -25000,
            "pretax_total_twd": 761905,
            "total_amount_twd": 800000,
        },
    })
    bill.ocr_lines = [
        OcrLine("計費期間：", 96, 1695, 125, 31, "user_info"),
        OcrLine("109/12/01至109/12/31", 342, 1697, 287, 30, "user_info"),
        OcrLine("稅前應繳總金額", 710, 891, 186, 28, "charge_info"),
        OcrLine("761905.0元", 1020, 893, 172, 31, "charge_info"),
    ]
    return bill


class FakeExtractor:
    cross_checks = False
    uses_paddle = False

    def extract(self, path, page_number):
        return official_sample_bill()


def png_bytes(size=(1654, 2339)):
    buffer = io.BytesIO()
    Image.new("RGB", size, "white").save(buffer, format="PNG")
    return buffer.getvalue()


def landlord_statement():
    """A landlord's allocation for a building with one meter (not a Taipower bill)."""
    statement = BillResult({"document_type": "unknown", "quality": {"warnings": []}})
    cells = [
        ("電費分攤明細表", 400, 80, 260),
        ("計費期間：113/01/01 ~ 113/01/31", 100, 140, 520),
        ("全棟總度數", 100, 200, 160), ("85,000 度", 300, 200, 140),
        ("承租戶", 100, 280, 100), ("分攤度數", 380, 280, 130), ("未稅金額", 580, 280, 130), ("含稅金額", 920, 280, 130),
        ("乙公司", 100, 330, 100), ("20,000", 390, 330, 100), ("60,000", 590, 330, 100), ("63,000", 930, 330, 100),
        ("甲公司", 100, 380, 100), ("15,000", 390, 380, 100), ("45,000", 590, 380, 100), ("47,250", 930, 380, 100),
    ]
    statement.ocr_lines = [OcrLine(text, x, y, width, 30, None) for text, x, y, width in cells]
    return statement


def energy_saving_page():
    """A bill's back page: the 節電資訊 table with this period's usage."""
    page = BillResult({
        "document_type": "high_voltage_bill",
        "quality": {"warnings": []},
        "energy_saving_info": {"comparison_rows": [
            {"period": "current", "days": 31, "energy_kwh": 15000, "daily_average_kwh": 483.87},
            {"period": "same_period_last_year", "days": 31, "energy_kwh": 20000, "daily_average_kwh": 645.16},
        ]},
    })
    page.ocr_lines = [
        OcrLine("用電日數", 1062, 112, 105, 32, "energy_saving_info"),
        OcrLine("度數", 1261, 114, 57, 31, "energy_saving_info"),
        OcrLine("本期", 860, 147, 59, 34, "energy_saving_info"),
        OcrLine("31", 1086, 150, 37, 30, "energy_saving_info"),
        OcrLine("15000", 1254, 150, 78, 29, "energy_saving_info"),
        OcrLine("483.87", 1438, 150, 85, 29, "energy_saving_info"),
        OcrLine("去年同期", 862, 186, 105, 29, "energy_saving_info"),
        OcrLine("20000", 1254, 186, 77, 29, "energy_saving_info"),
    ]
    return page


def three_page_pdf(pages=3):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.settings = Settings(data_dir=root / "data", output_dir=root / "out", web_dist=root / "no-web")
        self.app = create_app(self.settings, extractor_factory=FakeExtractor)
        self.client = TestClient(self.app)
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.app.state.store._db.close()
        self.temp.cleanup()

    def upload(self, user, *names):
        files = [("files", (name, png_bytes() if name.endswith(".png") else b"text", "application/octet-stream")) for name in names]
        response = self.client.post("/api/uploads", data={"user": user}, files=files)
        self.app.state.worker.wait_idle(30)
        return response.json()

    def first_item(self, user):
        return self.client.get("/api/items", params={"user": user}).json()[0]

    def test_batch_upload_processes_bills_and_rejects_other_files(self):
        result = self.upload("Amy", "jan.png", "feb.png", "notes.txt")
        self.assertEqual([item["filename"] for item in result["accepted"]], ["jan.png", "feb.png"])
        self.assertEqual(result["rejected"], [{"filename": "notes.txt", "reason": "unsupported file type"}])
        uploads = self.client.get("/api/uploads", params={"user": "Amy"}).json()
        self.assertEqual({upload["status"] for upload in uploads}, {"done"})
        items = self.client.get("/api/items", params={"user": "Amy"}).json()
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0]["billing_period"], "109/12/01~109/12/31")

    def test_evidence_marks_cover_billing_period_and_pretax_amount(self):
        self.upload("Amy", "jan.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        labels = {mark["text"]: mark for mark in detail["marks"] if mark["type"] == "label"}
        self.assertEqual(set(labels), {"Billing period", "Grey Electricity Cost"})
        highlights = [mark for mark in detail["marks"] if mark["type"] == "highlight"]
        pretax = next(mark for mark in highlights if mark["group"] == labels["Grey Electricity Cost"]["group"])
        # The highlight covers the label and the amount on the 稅前應繳總金額 row.
        self.assertLessEqual(pretax["x"], 710)
        self.assertGreaterEqual(pretax["x"] + pretax["w"], 1020 + 172)
        self.assertEqual([note for note in detail["mark_notes"] if not note.startswith("Electricity Usage")], [])

    def test_residential_bill_marks_the_total_and_says_so(self):
        bill = official_sample_bill()
        del bill["charge_info"]["pretax_total_twd"]
        bill.ocr_lines = [line for line in bill.ocr_lines if "稅前" not in line.text and "761905" not in line.text]
        bill.ocr_lines.append(OcrLine("應繳總金額", 710, 960, 152, 37, "charge_info"))
        bill.ocr_lines.append(OcrLine("800000元", 1014, 965, 177, 34, "charge_info"))
        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: bill):
            self.upload("Amy", "home.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        self.assertIn("Grey Electricity Cost: no pre-tax line on this bill, marked 應繳總金額 instead", detail["mark_notes"])

    def test_edits_recompute_other_adjustments_and_marks_are_saved(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        marks = self.client.get(f"/api/items/{item_id}").json()["marks"]
        marks[0]["x"] += 25
        response = self.client.put(f"/api/items/{item_id}", json={"user": "Amy", "row": {"流動電費": 125000}, "marks": marks})
        detail = response.json()
        self.assertEqual(detail["row"]["流動電費"], 125000)
        self.assertEqual(detail["row"]["其他調整費用"], 0)  # 800000 - 700000 - 125000 + 25000
        self.assertEqual(detail["marks"][0]["x"], marks[0]["x"])

    def test_evidence_downloads(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        png = self.client.get(f"/api/items/{item_id}/evidence.png")
        self.assertEqual(png.headers["content-type"], "image/png")
        marked = Image.open(io.BytesIO(png.content)).convert("RGB")
        self.assertNotEqual(marked.getpixel((800, 900)), (255, 255, 255))  # highlighted
        pdf = self.client.get(f"/api/items/{item_id}/evidence.pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_submit_writes_evidence_and_both_workbooks(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        result = self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy", "row": {"基本電費": 700100}}).json()
        self.assertEqual(result["status"], "submitted")
        self.assertEqual(result["locked_workbooks"], [])
        self.assertTrue(Path(result["evidence_file"]).read_bytes().startswith(b"%PDF"))

        report = load_workbook(self.settings.report_workbook)["Bills"]
        header = [cell.value for cell in report[1]]
        row = dict(zip(header, [cell.value for cell in report[2]]))
        self.assertEqual((row["基本電費"], row["上傳者"], row["狀態"]), (700100, "Amy", "submitted"))
        self.assertTrue(row["需人工檢查"].startswith("已人工確認"))
        self.assertTrue(row["證據檔"].endswith(".pdf"))

        fields = load_workbook(self.settings.master_workbook)["Fields"]
        header = [cell.value for cell in fields[1]]
        rows = [dict(zip(header, [cell.value for cell in line])) for line in fields.iter_rows(min_row=2)]
        basic = next(row for row in rows if row["key"] == "basic_charge_twd")
        self.assertEqual((basic["value"], basic["status"], basic["cross_check_value"]), (700100, "reviewer_corrected", 700000))

    def test_submitted_bill_is_read_only_until_reopened(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"})
        self.assertEqual(self.client.put(f"/api/items/{item_id}", json={"user": "Amy", "row": {"基本電費": 1}}).status_code, 409)
        reopened = self.client.post(f"/api/items/{item_id}/reopen", json={"user": "Amy"}).json()
        self.assertEqual(reopened["status"], "draft")
        self.assertEqual(load_workbook(self.settings.report_workbook)["Bills"].max_row, 1)  # header only

    def test_export_contains_only_the_users_own_bills(self):
        self.upload("Amy", "jan.png")
        self.upload("Ben", "feb.png")
        content = self.client.get("/api/export.xlsx", params={"user": "Amy"}).content
        sheet = load_workbook(io.BytesIO(content)).active
        header = [cell.value for cell in sheet[1]]
        rows = [dict(zip(header, [cell.value for cell in line])) for line in sheet.iter_rows(min_row=2)]
        self.assertEqual([(row["上傳者"], row["來源檔案"]) for row in rows], [("Amy", "jan.png")])

    def test_all_evidence_pdfs_download_as_one_zip(self):
        import zipfile

        self.upload("Amy", "jan.png", "feb.png")
        self.upload("Ben", "mar.png")
        first = self.first_item("Amy")["id"]
        submitted = self.client.post(f"/api/items/{first}/submit", json={"user": "Amy"}).json()

        response = self.client.get("/api/evidence.zip", params={"user": "Amy"})
        self.assertEqual(response.headers["content-type"], "application/zip")
        archive = zipfile.ZipFile(io.BytesIO(response.content))
        names = sorted(archive.namelist())
        self.assertEqual(len(names), 2)  # Amy's two bills, not Ben's
        self.assertEqual([name.split("/")[0] for name in names], ["draft", "submitted"])
        saved = Path(submitted["evidence_file"])
        self.assertEqual(names, ["draft/feb.pdf", "submitted/jan.pdf"])  # the uploaded files' names
        self.assertEqual(archive.read("submitted/jan.pdf"), saved.read_bytes())
        self.assertTrue(all(archive.read(name).startswith(b"%PDF") for name in names))

        only_submitted = zipfile.ZipFile(io.BytesIO(self.client.get("/api/evidence.zip", params={"user": "Amy", "status": "submitted"}).content))
        self.assertEqual(only_submitted.namelist(), ["submitted/jan.pdf"])
        self.assertEqual(self.client.get("/api/evidence.zip", params={"user": "Nobody"}).status_code, 404)

    def test_workbook_open_in_excel_is_reported(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        with mock.patch.object(reports, "save_tables", return_value=False):
            result = self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"}).json()
        self.assertIn("taipower_bills_17.xlsx", result["locked_workbooks"])

    def test_value_found_only_by_the_cross_check_engine_is_marked_exactly(self):
        # Paddle missed the period; Windows OCR read it, so its line is used.
        bill = official_sample_bill()
        bill.ocr_lines = [line for line in bill.ocr_lines if "109/12" not in line.text and "計費期間" not in line.text]
        bill["blocks"] = {"user_info": {"bounds": {"x": 96, "y": 634, "width": 559, "height": 1188}}}
        windows = official_sample_bill()
        FakeExtractor.cross_checks = True
        try:
            with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: bill), \
                 mock.patch.object(FakeExtractor, "extract_cross_check", lambda self, path, page: windows, create=True):
                self.upload("Amy", "jan.png")
        finally:
            FakeExtractor.cross_checks = False
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        period = detail["locations"]["Billing Period 計費期間"]
        self.assertTrue(period["exact"])
        self.assertLess(period["box"][3], 80)  # one line, not the whole block
        self.assertEqual([note for note in detail["mark_notes"] if not note.startswith("Electricity Usage")], [])

    def test_block_only_location_gets_a_small_highlight_to_drag(self):
        bill = official_sample_bill()
        bill.ocr_lines = [line for line in bill.ocr_lines if "計費期間" not in line.text and "109/12" not in line.text]
        bill["blocks"] = {"user_info": {"bounds": {"x": 96, "y": 634, "width": 559, "height": 1188}}}
        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: bill):
            self.upload("Amy", "jan.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        labels = {mark["text"]: mark for mark in detail["marks"] if mark["type"] == "label"}
        period = next(m for m in detail["marks"] if m["type"] == "highlight" and m["group"] == labels["Billing period"]["group"])
        self.assertLess(period["h"], 100)  # not the 1188 px block
        self.assertIn("Billing period: exact line not found, drag the highlight onto Billing Period 計費期間", detail["mark_notes"])

    def test_image_with_nothing_to_find_is_filled_in_by_hand(self):
        empty = BillResult({"document_type": "unknown", "quality": {"warnings": ["legacy_virtual_blocks"]}, "charge_info": {}})
        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: empty):
            self.upload("Amy", "photo.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        self.assertEqual(detail["kind"], "other")
        self.assertEqual(detail["marks"], [])
        self.assertTrue(all("not found" in issue["en"] for issue in detail["issues"]))
        # Every column stays open for the reviewer to fill in.
        columns = [column["key"] for column in self.client.get("/api/columns").json()]
        self.assertEqual(set(detail["row"]), set(columns))

    def submit_all(self, user):
        store = self.app.state.store
        for page in store.pages(user=user, bills_only=True):
            store.update_page(page["id"], status="submitted", submitted_at="2026-10-14T09:00:00", submitted_by=user)

    def test_simultaneous_workbook_writes_do_not_clash(self):
        import threading

        self.upload("Amy", *[f"bill{n}.png" for n in range(6)])
        self.submit_all("Amy")
        errors = []

        def write():
            try:
                self.app.state.writer.write()
            except Exception as exc:  # pragma: no cover - the failure being tested
                errors.append(exc)

        threads = [threading.Thread(target=write) for _ in range(8)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        self.assertEqual(errors, [])
        self.assertEqual(load_workbook(self.settings.report_workbook)["Bills"].max_row, 7)  # header + 6
        self.assertEqual([p.name for p in self.settings.output_dir.iterdir() if p.name.startswith("~tmp")], [])

    def test_workbook_open_in_excel_is_written_by_the_background_retry(self):
        import time

        writer = reports.WorkbookWriter(self.app.state.store, self.settings.output_dir, retry_seconds=0.1)
        writer.start()
        self.addCleanup(writer.stop)
        self.upload("Amy", "jan.png")
        self.submit_all("Amy")
        real_save = reports.save_tables
        with mock.patch.object(reports, "save_tables", return_value=False):
            self.assertEqual(writer.write(), ["taipower_bills_17.xlsx"])
            self.assertTrue(writer.pending)
        # The file is closed again: the retry loop replaces it without a submit.
        self.assertIs(reports.save_tables, real_save)
        deadline = time.time() + 5
        while writer.pending and time.time() < deadline:
            time.sleep(0.05)
        self.assertFalse(writer.pending)
        self.assertEqual(load_workbook(self.settings.report_workbook)["Bills"].max_row, 2)

    def test_both_workbooks_record_uploader_and_submitter(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"})

        report = load_workbook(self.settings.report_workbook)["Bills"]
        header = [cell.value for cell in report[1]]
        row = dict(zip(header, [cell.value for cell in report[2]]))
        self.assertEqual((row["上傳者"], row["送出者"]), ("Amy", "Amy"))
        self.assertTrue(row["送出時間"])

        pages = load_workbook(self.settings.master_workbook)["Pages"]
        header = [cell.value for cell in pages[1]]
        page = dict(zip(header, [cell.value for cell in pages[2]]))
        self.assertEqual((page["uploaded_by"], page["submitted_by"]), ("Amy", "Amy"))
        self.assertEqual(page["submitted_at"], row["送出時間"])

    def test_reopen_moves_the_submitted_pdf_to_the_archive(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        saved = Path(self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"}).json()["evidence_file"])
        content = saved.read_bytes()

        reopened = self.client.post(f"/api/items/{item_id}/reopen", json={"user": "Amy"}).json()
        archived = Path(reopened["archived_evidence"])
        self.assertFalse(saved.exists())
        self.assertEqual(archived.parent, self.settings.evidence_dir / "_reopened" / "Amy")
        self.assertTrue(archived.name.startswith(saved.stem))
        self.assertEqual(archived.read_bytes(), content)

        resubmitted = self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"}).json()
        self.assertTrue(Path(resubmitted["evidence_file"]).is_file())
        # Reopening a draft has nothing to archive.
        self.client.post(f"/api/items/{item_id}/reopen", json={"user": "Amy"})
        self.assertIsNone(self.client.post(f"/api/items/{item_id}/reopen", json={"user": "Amy"}).json()["archived_evidence"])

    def test_user_name_cannot_escape_the_evidence_folder(self):
        self.upload("../Amy", "jan.png")
        item_id = self.first_item("../Amy")["id"]
        saved = Path(self.client.post(f"/api/items/{item_id}/submit", json={"user": "../Amy"}).json()["evidence_file"])
        self.assertEqual(saved.parent.parent, self.settings.evidence_dir)

    def uploads_of(self, user):
        return {upload["filename"]: upload for upload in self.client.get("/api/uploads", params={"user": user}).json()}

    def test_pause_stops_after_the_current_page_and_resume_reads_the_rest(self):
        import threading
        import time

        started, proceed, read = threading.Event(), threading.Event(), []

        def extract(self_, path, page):
            read.append((Path(path).suffix, page))
            if page == 1:
                started.set()
                proceed.wait(10)
            return official_sample_bill()

        with mock.patch.object(FakeExtractor, "extract", extract):
            files = [("files", ("three.pdf", three_page_pdf(), "application/pdf")), ("files", ("later.png", png_bytes(), "image/png"))]
            self.client.post("/api/uploads", data={"user": "Amy"}, files=files)
            self.assertTrue(started.wait(10))
            self.assertEqual(self.client.post("/api/uploads/pause", json={"user": "Amy"}).json(), {"paused": 2})
            proceed.set()
            deadline = time.time() + 10
            while self.uploads_of("Amy")["three.pdf"]["status"] != "paused" and time.time() < deadline:
                time.sleep(0.05)
            self.app.state.worker.wait_idle(10)
            uploads = self.uploads_of("Amy")
            self.assertEqual((uploads["three.pdf"]["status"], uploads["three.pdf"]["pages_done"]), ("paused", 1))
            self.assertEqual(uploads["later.png"]["status"], "paused")
            self.assertEqual(read, [(".pdf", 1)])

            self.assertEqual(self.client.post("/api/uploads/resume", json={"user": "Amy"}).json(), {"resumed": 2})
            self.app.state.worker.wait_idle(10)
        uploads = self.uploads_of("Amy")
        self.assertEqual((uploads["three.pdf"]["status"], uploads["three.pdf"]["pages_done"]), ("done", 3))
        self.assertEqual(uploads["later.png"]["status"], "done")
        # Page 1 was not read again.
        self.assertEqual(read, [(".pdf", 1), (".pdf", 2), (".pdf", 3), (".png", None)])
        self.assertEqual(len(self.client.get("/api/items", params={"user": "Amy"}).json()), 4)

    def test_pausing_does_not_touch_other_users_uploads(self):
        with mock.patch.object(self.app.state.worker, "submit"):  # keep everything queued
            self.upload("Amy", "a.png")
            self.client.post("/api/uploads", data={"user": "Ben"}, files=[("files", ("b.png", png_bytes(), "image/png"))])
        self.client.post("/api/uploads/pause", json={"user": "Amy"})
        self.assertEqual(self.uploads_of("Amy")["a.png"]["status"], "paused")
        self.assertEqual(self.uploads_of("Ben")["b.png"]["status"], "queued")

    def test_failed_upload_can_be_retried_or_removed(self):
        def broken(self_, path, page):
            raise RuntimeError("cannot read")

        with mock.patch.object(FakeExtractor, "extract", broken):
            self.upload("Amy", "bad.png", "worse.png")
        uploads = self.uploads_of("Amy")
        self.assertEqual({u["status"] for u in uploads.values()}, {"failed"})

        retried = self.client.post(f"/api/uploads/{uploads['bad.png']['id']}/retry", json={"user": "Amy"})
        self.assertEqual(retried.json()["status"], "queued")
        self.app.state.worker.wait_idle(10)
        self.assertEqual(self.uploads_of("Amy")["bad.png"]["status"], "done")

        stored = Path(self.app.state.store.upload(uploads["worse.png"]["id"])["path"])
        self.assertEqual(self.client.delete(f"/api/uploads/{uploads['worse.png']['id']}", params={"user": "Ben"}).status_code, 404)
        self.assertEqual(self.client.delete(f"/api/uploads/{uploads['worse.png']['id']}", params={"user": "Amy"}).status_code, 200)
        self.assertNotIn("worse.png", self.uploads_of("Amy"))
        self.assertFalse(stored.exists())
        # A finished upload stays with its bills.
        self.assertEqual(self.client.delete(f"/api/uploads/{uploads['bad.png']['id']}", params={"user": "Amy"}).status_code, 409)
        self.assertEqual(self.client.post(f"/api/uploads/{uploads['bad.png']['id']}/retry", json={"user": "Amy"}).status_code, 409)

    def upload_two_page_bill(self, user="Amy"):
        def extract(self_, path, page):
            return official_sample_bill() if page == 1 else energy_saving_page()

        with mock.patch.object(FakeExtractor, "extract", extract):
            self.client.post("/api/uploads", data={"user": user}, files=[("files", ("two.pdf", three_page_pdf(2), "application/pdf"))])
            self.app.state.worker.wait_idle(10)
        return self.client.get(f"/api/items/{self.first_item(user)['id']}").json()

    def test_current_usage_on_the_back_page_is_marked(self):
        from pypdf import PdfReader

        detail = self.upload_two_page_bill()
        self.assertEqual([(page["page"], page["bill"]) for page in detail["pages"]], [(1, True), (2, False)])
        usage = [mark for mark in detail["marks"] if mark.get("page") == 2]
        label = next(mark for mark in usage if mark["type"] == "label")
        highlight = next(mark for mark in usage if mark["type"] == "highlight")
        self.assertEqual(label["text"], "Electricity Usage (kWh)")
        # From the 本期 label to the 15000 under 度數, on that row only.
        self.assertLessEqual(highlight["x"], 860)
        self.assertGreaterEqual(highlight["x"] + highlight["w"], 1254 + 78)
        self.assertLess(highlight["h"], 60)
        self.assertEqual(detail["mark_notes"], [])

        item_id = detail["id"]
        self.assertEqual(self.client.get(f"/api/items/{item_id}/image", params={"page": 2}).status_code, 200)
        self.assertEqual(self.client.get(f"/api/items/{item_id}/image", params={"page": 3}).status_code, 404)
        self.assertEqual(self.client.get(f"/api/items/{item_id}/evidence.png", params={"page": 2}).status_code, 200)
        self.assertEqual(len(PdfReader(io.BytesIO(self.client.get(f"/api/items/{item_id}/evidence.pdf").content)).pages), 2)

        # The back page stays in the evidence even without marks.
        bill_marks = [mark for mark in detail["marks"] if "page" not in mark]
        self.client.put(f"/api/items/{item_id}", json={"user": "Amy", "marks": bill_marks})
        self.assertEqual(len(PdfReader(io.BytesIO(self.client.get(f"/api/items/{item_id}/evidence.pdf").content)).pages), 2)
        # Reset puts the usage mark back.
        reset = self.client.post(f"/api/items/{item_id}/reset-marks").json()
        self.assertEqual(sum(mark.get("page") == 2 for mark in reset["marks"]), 2)

        saved = Path(self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"}).json()["evidence_file"])
        self.assertEqual(len(PdfReader(str(saved)).pages), 2)

    def test_bill_without_usage_says_so(self):
        self.upload("Amy", "jan.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        self.assertEqual(len(detail["pages"]), 1)
        self.assertTrue(any(note.startswith("Electricity Usage (kWh): 本期度數") for note in detail["mark_notes"]))

    def test_usage_is_never_marked_on_the_bill_page(self):
        # A 本期 row on the bill page itself is not 節電資訊's: leave it alone.
        bill = official_sample_bill()
        bill.ocr_lines += energy_saving_page().ocr_lines
        bill["energy_saving_info"] = energy_saving_page()["energy_saving_info"]
        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: bill):
            self.upload("Amy", "front.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        self.assertFalse(any(mark.get("text") == "Electricity Usage (kWh)" for mark in detail["marks"]))
        self.assertTrue(any("not found on the back page" in note for note in detail["mark_notes"]))

    def test_remove_all_clears_the_queue_except_the_file_being_read(self):
        def broken(self_, path, page):
            raise RuntimeError("cannot read")

        with mock.patch.object(FakeExtractor, "extract", broken):
            self.upload("Amy", "bad.png")
        with mock.patch.object(self.app.state.worker, "submit"):  # stay queued
            self.upload("Amy", "a.png", "b.png")
        self.client.post("/api/uploads/pause", json={"user": "Amy"})
        self.upload("Ben", "c.png")
        self.assertEqual(self.client.post("/api/uploads/remove", json={"user": "Amy"}).json(), {"removed": 3})
        self.assertEqual(self.uploads_of("Amy"), {})
        self.assertEqual(list(self.uploads_of("Ben")), ["c.png"])

    def test_deleting_bills_removes_drafts_and_clears_submitted_from_the_list(self):
        self.upload("Amy", "jan.png", "feb.png")
        self.upload_two_page_bill()
        items = {item["filename"]: item["id"] for item in self.client.get("/api/items", params={"user": "Amy"}).json()}
        self.client.post(f"/api/items/{items['jan.png']}/submit", json={"user": "Amy"})
        jan_upload = self.app.state.store.page(items["jan.png"])["upload_id"]

        result = self.client.post("/api/items/delete", json={"user": "Amy", "ids": list(items.values())}).json()
        self.assertEqual(sorted(result["deleted"]), sorted([items["feb.png"], items["two.pdf"]]))
        self.assertEqual(result["hidden"], [items["jan.png"]])
        self.assertEqual(self.client.get("/api/items", params={"user": "Amy"}).json(), [])
        # Deleted drafts take their upload (and back pages) with them.
        self.assertEqual(list(self.uploads_of("Amy")), ["jan.png"])
        self.assertEqual([p["upload_id"] for p in self.app.state.store.pages(user="Amy")], [jan_upload])
        # The submitted bill is still in the Excel file and keeps its evidence.
        report = load_workbook(self.settings.report_workbook)["Bills"]
        self.assertEqual(report.max_row, 2)
        self.app.state.writer.write()
        self.assertEqual(load_workbook(self.settings.report_workbook)["Bills"].max_row, 2)
        self.assertEqual(len(list(self.settings.evidence_dir.rglob("*.pdf"))), 1)

    def test_cannot_delete_other_users_bills_or_bills_being_read(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        self.assertEqual(self.client.post("/api/items/delete", json={"user": "Ben", "ids": [item_id]}).json(), {"deleted": [], "hidden": []})
        self.app.state.store.set_upload(self.app.state.store.page(item_id)["upload_id"], status="processing")
        self.assertEqual(self.client.post("/api/items/delete", json={"user": "Amy", "ids": [item_id]}).status_code, 409)

    def test_chinese_file_name_downloads(self):
        from urllib.parse import unquote

        self.upload("Amy", "台電帳單.png")
        response = self.client.get(f"/api/items/{self.first_item('Amy')['id']}/evidence.pdf")
        self.assertEqual(response.status_code, 200)
        header = response.headers["content-disposition"]
        self.assertIn('filename="', header)
        self.assertEqual(unquote(header.split("filename*=UTF-8''")[1]), "台電帳單.pdf")

    def test_bills_from_one_file_get_numbered_names_in_the_zip(self):
        import zipfile

        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: official_sample_bill()):
            self.client.post("/api/uploads", data={"user": "Amy"}, files=[("files", ("台電.pdf", three_page_pdf(2), "application/pdf"))])
            self.app.state.worker.wait_idle(10)
        archive = zipfile.ZipFile(io.BytesIO(self.client.get("/api/evidence.zip", params={"user": "Amy"}).content))
        self.assertEqual(sorted(archive.namelist()), ["draft/台電 (2).pdf", "draft/台電.pdf"])

    def test_saved_evidence_pdf_keeps_the_uploaded_name(self):
        self.upload("Amy", "jan.png")
        self.upload("Amy", "jan.png")  # the same file again
        first, second = (item["id"] for item in self.client.get("/api/items", params={"user": "Amy"}).json())
        submit = lambda item_id: Path(self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"}).json()["evidence_file"])
        self.assertEqual(submit(first).name, "jan.pdf")
        self.assertEqual(submit(second).name, "jan (2).pdf")
        report = load_workbook(self.settings.report_workbook)["Bills"]
        header = [cell.value for cell in report[1]]
        self.assertEqual(sorted(row[header.index("證據檔")].replace("\\", "/") for row in report.iter_rows(min_row=2, values_only=True)),
                         ["evidence/Amy/jan (2).pdf", "evidence/Amy/jan.pdf"])
        # Reopened and sent again: the freed name is used again.
        self.client.post(f"/api/items/{first}/reopen", json={"user": "Amy"})
        self.assertEqual(submit(first).name, "jan.pdf")

    def local_client(self):
        # Requests from the PC running the server (the default client is remote).
        return TestClient(self.app, client=("127.0.0.1", 50000))

    def test_settings_can_be_changed_only_on_the_server_pc(self):
        view = self.client.get("/api/settings").json()
        self.assertFalse(view["can_edit"])
        self.assertEqual(view["values"]["host"], "127.0.0.1")
        self.assertEqual(self.client.put("/api/settings", json={"values": {"ocr_backend": "windows"}}).status_code, 403)
        self.assertTrue(self.local_client().get("/api/settings").json()["can_edit"])

    def test_settings_apply_at_once_and_are_saved(self):
        local = self.local_client()
        result = local.put("/api/settings", json={"values": {"ocr_backend": "paddleocr", "workbook_retry_seconds": 30, "port": 9000}}).json()
        self.assertEqual(result["changed"], ["ocr_backend", "port", "workbook_retry_seconds"])
        self.assertTrue(result["restart"])  # the port is used at the next start
        self.assertTrue(self.app.state.worker._reload)
        self.assertEqual(self.app.state.writer.retry_seconds, 30)
        saved = Settings.load(self.settings.data_dir)
        self.assertEqual((saved.ocr_backend, saved.port, saved.workbook_retry_seconds), ("paddleocr", 9000, 30))

    def test_invalid_settings_change_nothing(self):
        local = self.local_client()
        response = local.put("/api/settings", json={"values": {"ocr_backend": "windows", "port": 80}})
        self.assertEqual(response.status_code, 400)
        self.assertEqual((self.settings.ocr_backend, self.settings.port), ("crosscheck", 8765))
        self.assertFalse(self.settings.config_file.exists())

    def test_new_output_folder_keeps_earlier_evidence(self):
        import zipfile

        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        old_pdf = Path(self.client.post(f"/api/items/{item_id}/submit", json={"user": "Amy"}).json()["evidence_file"])
        new_output = Path(self.temp.name) / "sharepoint"
        result = self.local_client().put("/api/settings", json={"values": {"output_dir": str(new_output)}}).json()
        self.assertEqual(result["locked_workbooks"], [])
        # The workbooks are rebuilt in the new folder and still point at the PDF.
        report = load_workbook(new_output / "taipower_bills_17.xlsx")["Bills"]
        header = [cell.value for cell in report[1]]
        self.assertEqual(Path(report.cell(2, header.index("證據檔") + 1).value), old_pdf)
        archive = zipfile.ZipFile(io.BytesIO(self.client.get("/api/evidence.zip", params={"user": "Amy"}).content))
        self.assertEqual(archive.read("submitted/jan.pdf"), old_pdf.read_bytes())
        # New submissions go to the new folder.
        self.upload("Amy", "feb.png")
        feb = next(item["id"] for item in self.client.get("/api/items", params={"user": "Amy"}).json() if item["filename"] == "feb.png")
        new_pdf = Path(self.client.post(f"/api/items/{feb}/submit", json={"user": "Amy"}).json()["evidence_file"])
        self.assertEqual(new_pdf.parent, new_output / "evidence" / "Amy")

    def test_marks_use_their_own_colour_or_the_default(self):
        self.upload("Amy", "jan.png")
        item_id = self.first_item("Amy")["id"]
        self.assertEqual(self.client.get(f"/api/items/{item_id}").json()["colors"], {"highlight": "#ffeb3b", "label": "#d32f2f"})
        marks = [
            {"id": "a", "type": "highlight", "x": 100, "y": 100, "w": 200, "h": 60, "color": "#0000ff"},
            {"id": "b", "type": "highlight", "x": 100, "y": 300, "w": 200, "h": 60},
        ]
        self.client.put(f"/api/items/{item_id}", json={"user": "Amy", "marks": marks})
        # The default changes on the settings page; marks without a colour follow it.
        self.local_client().put("/api/settings", json={"values": {"highlight_color": "#00ff00"}})
        image = Image.open(io.BytesIO(self.client.get(f"/api/items/{item_id}/evidence.png").content)).convert("RGB")
        self.assertEqual(image.getpixel((200, 130)), (153, 153, 255))  # blue at 40% on white
        self.assertEqual(image.getpixel((200, 330)), (153, 255, 153))  # the new default, green
        bad = self.client.put(f"/api/items/{item_id}", json={"user": "Amy", "marks": [{**marks[0], "color": "blue"}]})
        self.assertEqual(bad.status_code, 422)
        self.assertEqual(self.local_client().put("/api/settings", json={"values": {"label_color": "red"}}).status_code, 400)

    def upload_statement(self, user="Amy", name="statement.png"):
        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: landlord_statement()):
            self.upload(user, name)
        return self.client.get(f"/api/items/{self.first_item(user)['id']}").json()

    def test_landlord_statement_values_are_found_by_keyword(self):
        detail = self.upload_statement()
        self.assertEqual(detail["kind"], "other")
        row = detail["row"]
        self.assertEqual(
            (row["Billing Period 計費期間"], row["用電度數"], row["稅前應繳總金額"], row["應繳總金額"]),
            ("113/01/01~113/01/31", 20000, 60000, 63000),  # no company name set: the first tenant row
        )
        self.assertEqual(set(detail["sources"].values()), {"keyword"})
        labels = {mark["text"]: mark for mark in detail["marks"] if mark["type"] == "label"}
        self.assertEqual(set(labels), {"Billing period", "Electricity Usage (kWh)", "Grey Electricity Cost"})
        self.assertTrue(all(mark.get("field") for mark in detail["marks"]))
        # Every found value waits for the reviewer's check.
        self.assertTrue(any("check it" in issue["en"] for issue in detail["issues"]))
        self.assertTrue(detail["pages"][0]["lines"])

    def test_company_name_in_settings_picks_the_companys_row(self):
        detail = self.upload_statement()
        self.local_client().put("/api/settings", json={"values": {"company_names": "甲公司\n甲科技"}})
        again = self.client.post(f"/api/items/{detail['id']}/redetect").json()
        self.assertEqual((again["row"]["用電度數"], again["row"]["稅前應繳總金額"]), (15000, 45000))
        self.assertEqual(self.app.state.settings.company_names, ["甲公司", "甲科技"])

    def test_document_values_and_their_sources_reach_both_workbooks(self):
        detail = self.upload_statement()
        row = {"用電度數": 15000, "稅前應繳總金額": 45000}
        self.client.post(f"/api/items/{detail['id']}/submit", json={"user": "Amy", "row": row, "sources": {"用電度數": "pick", "稅前應繳總金額": "manual"}})
        report = load_workbook(self.settings.report_workbook)["Bills"]
        header = [cell.value for cell in report[1]]
        line = dict(zip(header, [cell.value for cell in report[2]]))
        self.assertEqual((line["文件類型"], line["用電度數"], line["稅前應繳總金額"]), ("其他單據", 15000, 45000))
        self.assertIn("用電度數:圖上點選", line["欄位來源"])
        self.assertIn("未稅金額:手動輸入", line["欄位來源"])
        self.assertIn("計費期間:關鍵字", line["欄位來源"])
        # Notes about values the reviewer picked or typed are dropped.
        self.assertIn("計費期間以關鍵字", line["需人工檢查"])
        self.assertNotIn("用電度數以關鍵字", line["需人工檢查"])

        fields = load_workbook(self.settings.master_workbook)["Fields"]
        header = [cell.value for cell in fields[1]]
        rows = [dict(zip(header, [cell.value for cell in item])) for item in fields.iter_rows(min_row=2)]
        usage = next(item for item in rows if item["block"] == "document" and item["key"] == "energy_kwh")
        self.assertEqual((usage["value"], usage["status"], usage["cross_check_value"]), (15000, "picked_by_reviewer", 20000))
        pages = load_workbook(self.settings.master_workbook)["Pages"]
        header = [cell.value for cell in pages[1]]
        self.assertEqual(dict(zip(header, [cell.value for cell in pages[2]]))["document_type"], "other_document")

    def test_taipower_bill_usage_comes_from_the_back_page(self):
        detail = self.upload_two_page_bill()
        self.assertEqual(detail["kind"], "taipower")
        self.assertEqual(detail["row"]["用電度數"], 15000)
        report_sources = self.client.post(f"/api/items/{detail['id']}/submit", json={"user": "Amy"})
        self.assertEqual(report_sources.status_code, 200)
        report = load_workbook(self.settings.report_workbook)["Bills"]
        header = [cell.value for cell in report[1]]
        line = dict(zip(header, [cell.value for cell in report[2]]))
        self.assertEqual((line["文件類型"], line["用電度數"], line["欄位來源"]), ("台電帳單", 15000, "其他:台電帳單辨識"))

    def test_evidence_of_another_document_keeps_every_page(self):
        from pypdf import PdfReader

        def extract(self_, path, page):
            return landlord_statement() if page == 1 else BillResult({"document_type": "unknown", "quality": {"warnings": []}})

        with mock.patch.object(FakeExtractor, "extract", extract):
            self.client.post("/api/uploads", data={"user": "Amy"}, files=[("files", ("notice.pdf", three_page_pdf(3), "application/pdf"))])
            self.app.state.worker.wait_idle(10)
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        self.assertEqual(detail["kind"], "other")
        self.assertEqual(len(detail["pages"]), 3)
        self.assertFalse(any(mark.get("page") for mark in detail["marks"]))  # all marks on page 1
        content = self.client.get(f"/api/items/{detail['id']}/evidence.pdf").content
        self.assertEqual(len(PdfReader(io.BytesIO(content)).pages), 3)
        saved = Path(self.client.post(f"/api/items/{detail['id']}/submit", json={"user": "Amy"}).json()["evidence_file"])
        self.assertEqual(len(PdfReader(str(saved)).pages), 3)

    def test_no_page_is_left_out_of_a_bill(self):
        from pypdf import PdfReader

        cover = BillResult({"document_type": "unknown", "quality": {"warnings": []}})
        pages = {1: cover, 2: official_sample_bill(), 3: energy_saving_page(), 4: cover}

        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: pages[page]):
            self.client.post("/api/uploads", data={"user": "Amy"}, files=[("files", ("four.pdf", three_page_pdf(4), "application/pdf"))])
            self.app.state.worker.wait_idle(10)
        items = self.client.get("/api/items", params={"user": "Amy"}).json()
        self.assertEqual([(item["page"], item["kind"]) for item in items], [(2, "taipower")])
        detail = self.client.get(f"/api/items/{items[0]['id']}").json()
        # The cover before the bill and both pages after it belong to it, in order.
        self.assertEqual([(view["page"], view["bill"]) for view in detail["pages"]], [(1, False), (2, True), (3, False), (4, False)])
        self.assertEqual(self.client.get(f"/api/items/{detail['id']}/image", params={"page": 1}).status_code, 200)
        content = self.client.get(f"/api/items/{detail['id']}/evidence.pdf").content
        self.assertEqual(len(PdfReader(io.BytesIO(content)).pages), 4)

    def test_redetect_is_only_for_other_documents(self):
        self.upload("Amy", "jan.png")
        self.assertEqual(self.client.post(f"/api/items/{self.first_item('Amy')['id']}/redetect").status_code, 409)

    def test_warnings_can_be_cleared_by_hand(self):
        bill = official_sample_bill()
        del bill["charge_info"]["basic_charge_twd"]  # a warning; the usage note is a mark note
        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: bill):
            self.upload("Amy", "jan.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        self.assertTrue(detail["issues"] and detail["mark_notes"])
        kept = detail["issues"][1:]
        saved = self.client.put(f"/api/items/{detail['id']}", json={"user": "Amy", "issues": kept, "mark_notes": []}).json()
        self.assertEqual([issue["en"] for issue in saved["issues"]], [issue["en"] for issue in kept])
        self.assertEqual(saved["mark_notes"], [])
        self.assertEqual(saved["row"]["需人工檢查"], "；".join(issue["zh"] for issue in kept) or None)
        # Clearing never brings a warning back or adds one.
        again = self.client.put(f"/api/items/{detail['id']}", json={"user": "Amy", "issues": detail["issues"] + [{"en": "made up", "zh": "x"}]}).json()
        self.assertEqual(len(again["issues"]), len(kept))

    def test_each_user_keeps_their_own_tags(self):
        defaults = self.client.get("/api/preferences", params={"user": "Amy"}).json()
        self.assertEqual(defaults["mark_tags"], ["Billing period", "Grey Electricity Cost", "Electricity Usage (kWh)"])
        saved = self.client.put("/api/preferences", json={"user": "Amy", "mark_tags": [" Meter No. ", "Billing period", "Meter No.", ""]}).json()
        self.assertEqual(saved["mark_tags"], ["Meter No.", "Billing period"])
        self.assertEqual(self.client.get("/api/preferences", params={"user": "Ben"}).json()["mark_tags"], defaults["mark_tags"])
        reset = self.client.put("/api/preferences", json={"user": "Amy", "mark_tags": None}).json()
        self.assertEqual(reset["mark_tags"], defaults["mark_tags"])

    def save_sites(self, sites):
        return self.local_client().put("/api/sites", json={"sites": sites})

    def test_sites_are_set_on_the_server_pc_and_meters_compared_by_digits(self):
        self.assertEqual(self.client.put("/api/sites", json={"sites": []}).status_code, 403)
        saved = self.save_sites([
            {"name": "台北總部", "code": "TW-HQ", "meters": ["00-12-3456-78-9", "00123456789", " 10512918917 "]},
            {"name": "新竹廠", "meters": ["10518691508"]},
        ]).json()
        self.assertEqual([(site["name"], site["meters"]) for site in saved],
                         [("台北總部", ["00123456789", "10512918917"]), ("新竹廠", ["10518691508"])])
        clash = self.save_sites([{"name": "A", "meters": ["1"]}, {"name": "B", "meters": ["1"]}])
        self.assertEqual(clash.status_code, 400)
        self.assertEqual(self.save_sites([{"name": "A"}, {"name": "A"}]).status_code, 400)

    def test_a_bill_belongs_to_the_site_of_its_meter_unless_chosen_by_hand(self):
        sites = self.save_sites([{"name": "台北總部", "meters": ["00-12-3456-78-9"]}, {"name": "新竹廠", "meters": []}]).json()
        hq, plant = (site["id"] for site in sites)
        self.upload("Amy", "jan.png")
        item = self.first_item("Amy")
        self.assertEqual(item["site"], "台北總部")
        detail = self.client.get(f"/api/items/{item['id']}").json()
        self.assertEqual((detail["site"]["name"], detail["site"]["how"]), ("台北總部", "meter"))
        moved = self.client.put(f"/api/items/{item['id']}", json={"user": "Amy", "site_id": plant}).json()
        self.assertEqual((moved["site"]["name"], moved["site"]["how"]), ("新竹廠", "manual"))
        back = self.client.put(f"/api/items/{item['id']}", json={"user": "Amy", "site_id": None}).json()
        self.assertEqual(back["site"]["name"], "台北總部")
        # Saving other things leaves the choice alone; removing the site clears it.
        self.client.put(f"/api/items/{item['id']}", json={"user": "Amy", "site_id": plant})
        self.client.put(f"/api/items/{item['id']}", json={"user": "Amy", "row": {"基本電費": 700000}})
        self.assertEqual(self.client.get(f"/api/items/{item['id']}").json()["site"]["name"], "新竹廠")
        self.save_sites([{"id": hq, "name": "台北總部", "meters": ["00123456789"]}])
        self.assertEqual(self.client.get(f"/api/items/{item['id']}").json()["site"]["how"], "meter")
        self.assertEqual(self.client.put(f"/api/items/{item['id']}", json={"user": "Amy", "site_id": 999}).status_code, 400)

    def test_monthly_sheet_splits_each_bill_by_days(self):
        self.save_sites([{"name": "台北總部", "code": "TW-HQ", "meters": ["00123456789"]}])
        bill = official_sample_bill()
        bill["user_info"]["billing_period"] = {"start": "113/12/15", "end": "114/01/14"}
        bill["energy_saving_info"] = None
        with mock.patch.object(FakeExtractor, "extract", lambda self, path, page: bill):
            self.upload("Amy", "jan.png")
        detail = self.client.get(f"/api/items/{self.first_item('Amy')['id']}").json()
        self.assertEqual([(month["月份"], month["天數"]) for month in detail["months"]], [("2024-12", 17), ("2025-01", 14)])
        self.client.post(f"/api/items/{detail['id']}/submit", json={"user": "Amy"})

        sheet = load_workbook(self.settings.report_workbook)["Monthly"]
        header = [cell.value for cell in sheet[1]]
        rows = [dict(zip(header, line)) for line in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertEqual([(row["據點"], row["據點代碼"], row["電號"], row["月份"], row["天數"]) for row in rows],
                         [("台北總部", "TW-HQ", "00123456789", "2024-12", 17), ("台北總部", "TW-HQ", "00123456789", "2025-01", 14)])
        # The parts add up to the bill exactly.
        self.assertEqual(sum(row["應繳總金額"] for row in rows), 800000)
        self.assertEqual(sum(row["用電度數"] for row in rows), 8000 + 3000 + 3000 + 1000)
        self.assertEqual(sum(row["稅前應繳總金額"] for row in rows), 761905)

    def test_user_name_is_required(self):
        response = self.client.post("/api/uploads", data={"user": "  "}, files=[("files", ("a.png", png_bytes(), "image/png"))])
        self.assertEqual(response.status_code, 400)


class PerUserWorkbookTests(unittest.TestCase):
    """Copies of the app on several PCs, each writing its own workbooks."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.settings = Settings(data_dir=root / "data", output_dir=root / "out", web_dist=root / "no-web", workbooks="per_user")
        self.app = create_app(self.settings, extractor_factory=FakeExtractor)
        self.client = TestClient(self.app)
        self.client.__enter__()

    tearDown = ApiTests.tearDown
    upload = ApiTests.upload
    first_item = ApiTests.first_item

    def rows(self, path):
        sheet = load_workbook(path)["Bills"]
        header = [cell.value for cell in sheet[1]]
        return [dict(zip(header, [cell.value for cell in line])) for line in sheet.iter_rows(min_row=2)]

    def test_each_user_gets_their_own_workbooks(self):
        self.upload("Amy", "jan.png")
        self.upload("Ben", "feb.png")
        for user in ("Amy", "Ben"):
            item_id = self.first_item(user)["id"]
            self.assertEqual(self.client.post(f"/api/items/{item_id}/submit", json={"user": user}).json()["locked_workbooks"], [])

        out = self.settings.output_dir
        self.assertFalse((out / "taipower_bills_17.xlsx").exists())
        self.assertEqual([row["來源檔案"] for row in self.rows(out / "Amy" / "taipower_bills_17.xlsx")], ["jan.png"])
        self.assertEqual([row["來源檔案"] for row in self.rows(out / "Ben" / "taipower_bills_17.xlsx")], ["feb.png"])
        self.assertTrue((out / "Amy" / "taipower_bills.xlsx").is_file())
        self.assertEqual(self.client.get("/api/output", params={"user": "Amy"}).json()["folder"], str(out / "Amy"))

        # Reopening rewrites only that user's workbooks.
        self.client.post(f"/api/items/{self.first_item('Amy')['id']}/reopen", json={"user": "Amy"})
        self.assertEqual(self.rows(out / "Amy" / "taipower_bills_17.xlsx"), [])
        self.assertEqual(len(self.rows(out / "Ben" / "taipower_bills_17.xlsx")), 1)

    def test_locked_user_workbook_is_retried_for_that_user(self):
        import time

        writer = reports.WorkbookWriter(self.app.state.store, self.settings.output_dir, retry_seconds=0.1, per_user=True)
        writer.start()
        self.addCleanup(writer.stop)
        self.upload("Amy", "jan.png")
        store = self.app.state.store
        for page in store.pages(user="Amy", bills_only=True):
            store.update_page(page["id"], status="submitted", submitted_at="2026-10-14T09:00:00", submitted_by="Amy")
        with mock.patch.object(reports, "save_tables", return_value=False):
            self.assertEqual(writer.write("Amy"), ["taipower_bills_17.xlsx"])
        deadline = time.time() + 5
        while writer.pending and time.time() < deadline:
            time.sleep(0.05)
        self.assertFalse(writer.pending)
        self.assertEqual(len(self.rows(self.settings.output_dir / "Amy" / "taipower_bills_17.xlsx")), 1)



class SettingsLoadTests(unittest.TestCase):
    def test_config_file_then_command_line_then_environment(self):
        import json
        import os

        with tempfile.TemporaryDirectory() as temp:
            data = Path(temp)
            (data / "config.json").write_text(json.dumps({"port": 9000, "ocr_backend": "windows", "host": "0.0.0.0"}), encoding="utf-8")
            with mock.patch.dict(os.environ, {"TAIPOWER_OCR_BACKEND": "paddleocr"}):
                settings = Settings.load(data, defaults={"workbooks": "per_user"}, overrides={"port": 9100})
            self.assertEqual((settings.host, settings.port, settings.ocr_backend, settings.workbooks), ("0.0.0.0", 9100, "paddleocr", "per_user"))
            self.assertEqual(settings.locked, frozenset({"ocr_backend"}))
            with self.assertRaises(ValueError):
                settings.update({"ocr_backend": "windows"})
            # Saving leaves the locked value out of config.json.
            settings.update({"port": 9200})
            settings.save()
            saved = json.loads((data / "config.json").read_text(encoding="utf-8"))
            self.assertEqual((saved["port"], saved["ocr_backend"]), (9200, "windows"))


if __name__ == "__main__":
    unittest.main()
